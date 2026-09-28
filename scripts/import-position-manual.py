"""Import the approved DOCX titles and manual into the shared organisation data.

Usage: python scripts/import-position-manual.py TITLES.docx MANUAL.docx [--check]
Requires python-docx. Paragraph pairs below are reviewed source correspondences;
an altered document must be reviewed before updating this mapping.
"""

import argparse
import json
import re
import unicodedata
from pathlib import Path

from docx import Document


# Title paragraph, manual heading paragraph, existing organisation identity.
POSITIONS = [
    (3, 1002, "division-controle-engagements-paiements"),
    (6, 1016, None), (9, 1027, None),
    (13, 1088, "division-etudes-affaires-juridiques"),
    (17, 1137, None), (20, 1104, None), (23, 1120, None),
    (27, 1039, "division-gouvernance-communication-externe"),
    (31, 1071, None), (35, 1058, None),
    (39, 969, "division-risk-management"),
    (42, 981, None), (45, 992, None), (49, 1150, "audit-interne"),
    (53, 244, "pole-clients"), (57, 257, "division-relation-client"),
    (59, 299, None), (62, 288, None), (65, 278, None), (68, 267, None),
    (72, 307, "division-experience-client"), (73, 328, None), (77, 319, None),
    (80, 348, "pole-gestion-portefeuille"),
    (84, 363, "division-gestion-portefeuille-operations"),
    (85, 376, None), (88, 386, None),
    (91, 398, "division-gestion-portefeuille-gestion"),
    (94, 408, None), (97, 431, None), (100, 418, None), (103, 445, None),
    (107, 105, "pole-operations"), (110, 159, "division-concession-droits"),
    (113, 179, None), (116, 170, None),
    (120, 115, "division-constitution-droits"),
    (124, 126, None), (127, 135, None), (130, 147, None),
    (134, 197, "division-paiement-prestations"),
    (137, 216, None), (140, 206, None), (143, 225, None),
    (146, 810, "pole-pilotage"), (149, 914, "division-data-office"),
    (152, 936, None), (155, 925, None),
    (158, 867, "division-organisation-amoa"), (161, 902, None), (164, 884, None),
    (167, 827, "division-planification-controle-gestion"),
    (171, 855, None), (174, 843, None), (177, 949, None),
    (181, 461, "pole-ressources"), (184, 483, "division-achats-logistique"),
    (187, 497, None), (190, 524, None), (193, 511, None),
    (197, 539, "division-capital-humain"), (200, 554, None), (203, 569, None),
    (207, 590, "division-financiere-comptable"),
    (210, 608, None), (213, 620, None), (216, 631, None),
    (220, 655, "pole-systeme-information-transformation-digitale"),
    (223, 672, "division-etude-developpement"), (226, 696, None), (229, 683, None),
    (232, 711, "division-exploitation-systemes-reseaux"),
    (235, 741, None), (238, 724, None),
    (241, 758, "division-securite-information"), (244, 771, None),
    (247, 783, None), (250, 793, None),
]


def normalize(text):
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def flatten(node):
    yield node
    for key in ("children", "hiddenChildren", "secondaryChildren"):
        for child in node.get(key, []):
            yield from flatten(child)


def extract_manual(document):
    paragraphs = document.paragraphs
    sections = {}
    current = None
    field = None
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if not text or paragraph.style.name.lower().startswith("toc"):
            continue
        if text.startswith("Missions principales"):
            heading = index - 1
            while not paragraphs[heading].text.strip():
                heading -= 1
            title = paragraphs[heading].text.strip()
            assert normalize(title).startswith("chef"), (heading, title)
            current = {"title": title, "missions": [], "competences": []}
            sections[heading] = current
            field = "missions"
        elif normalize(text).startswith("competences"):
            field = "competences"
        elif normalize(text).startswith(("chef-", "pole-", "entites-rattachees")):
            current = None
            field = None
        elif current is not None and field:
            current[field].append(text)
    return sections


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("titles", type=Path)
    parser.add_argument("manual", type=Path)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    titles = Document(args.titles).paragraphs
    manual = extract_manual(Document(args.manual))
    root = Path(__file__).resolve().parents[1]
    folder = root / "data/rubriques/organisation-smi-culture/sous-rubriques"
    org_path = folder / "organigramme.json"
    postes_path = folder / "fiches-et-fonctions-de-postes.json"
    org_file = json.loads(org_path.read_text(encoding="utf-8"))
    postes_file = json.loads(postes_path.read_text(encoding="utf-8"))
    org = org_file["data"]["orgData"]
    nodes = list(flatten(org))
    by_id = {node["posteId"]: node for node in nodes if node.get("posteId")}
    # These existing unoccupied entities originally had no fiche identifier.
    for name, identity in [
        ("Pôle Ressources", "pole-ressources"),
        ("Division Planification et Contrôle de Gestion", "division-planification-controle-gestion"),
        ("Division Sécurité de l'information", "division-securite-information"),
    ]:
        node = next(node for node in nodes if node["name"] == name)
        node["posteId"] = identity
        by_id[identity] = node

    org["hasFiche"] = False
    org["functionTitle"] = titles[0].text.strip()
    by_id["secretariat-general"]["hasFiche"] = False
    positions = []
    current_division = None
    current_pole = None
    mapped_indices = {entry[0] for entry in POSITIONS} | {0}
    source_indices = {i for i, p in enumerate(titles) if p.text.strip()
                      and not p.text.strip().startswith(("Missions", "Compétences"))}
    assert source_indices == mapped_indices, "The title document changed; review the mapping."
    assert len({entry[1] for entry in POSITIONS}) == len(POSITIONS)

    for title_index, manual_index, existing_id in POSITIONS:
        title = titles[title_index].text.strip()
        section = manual[manual_index]
        assert section["missions"] and section["competences"], title
        identity = existing_id or normalize(section["title"])
        if existing_id:
            node = by_id[identity]
        else:
            # Innovation and digital transformation are services reporting to their poles.
            parent = current_pole if title_index in (177, 250) else current_division
            assert parent is not None, title
            child_key = "hiddenChildren" if parent is current_pole else "children"
            children = parent.setdefault(child_key, [])
            node = next((child for child in children if child.get("posteId") == identity), None)
            if node is None:
                node = {"name": title, "role": "Service", "posteId": identity}
                children.append(node)
        node["hasFiche"] = True
        node["functionTitle"] = title
        if existing_id and existing_id.startswith("pole-"):
            current_pole = node
            current_division = None
        elif existing_id and existing_id.startswith("division-"):
            current_division = node
        positions.append({
            "id": identity,
            "titre": title,
            "famille": current_pole["name"] if current_pole else "Entités rattachées à la Direction",
            "missions": section["missions"],
            "competences": section["competences"],
            "folderAliases": [section["title"]],
            "attachments": [],
            "source": {"titlesDocument": args.titles.name, "titleParagraph": title_index,
                       "manualDocument": args.manual.name, "manualParagraph": manual_index},
        })
    assert len({p["id"] for p in positions}) == len(positions)
    final_nodes = list(flatten(org))
    assert {n["posteId"] for n in final_nodes if n.get("hasFiche")} == {p["id"] for p in positions}
    postes_file["data"]["postesData"] = positions
    postes_file["data"]["orgGovPages"]["postes"]["description"] = (
        "Consulter les missions, les compétences et les pièces jointes de chaque fonction."
    )
    for path, value in [(org_path, org_file), (postes_path, postes_file)]:
        if args.check:
            assert json.loads(path.read_text(encoding="utf-8")) == value, f"Out of date: {path}"
        else:
            path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f'{len(positions)} fiches: titres, missions et competences verifies; {len(final_nodes)} noeuds.')
    unused = [s["title"] for i, s in manual.items() if i not in {p[1] for p in POSITIONS}]
    print("Manual entries outside the title reference:", unused)


if __name__ == "__main__":
    main()
