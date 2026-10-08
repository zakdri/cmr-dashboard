import json
import re
from docx import Document


def clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


document = Document(r"C:\Users\alami\Downloads\Partie gouvernance.docx")
paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
with open(r"D:\cmr-dashboard\data\rubriques\gouvernance\sous-rubriques\systeme-de-gouvernance.json", encoding="utf-8") as source:
    system = json.load(source)["data"]
with open(r"D:\cmr-dashboard\data\rubriques\gouvernance\sous-rubriques\missions-et-valeurs.json", encoding="utf-8") as source:
    missions_data = json.load(source)["data"]

differences = []


def compare(label, site_value, document_value):
    if clean(site_value) != clean(document_value):
        differences.append((label, clean(site_value), clean(document_value)))


board = system["governanceBoard"]
compare("Conseil - introduction", board["introduction"], paragraphs[4])
compare("Conseil - description", board["description"], f"{paragraphs[7]} {paragraphs[8]}")
compare("Conseil - introduction missions", board["missionsIntroduction"], paragraphs[10])
for index, paragraph_index in enumerate(range(11, 17)):
    compare(f"Conseil - mission {index + 1}", board["missions"][index], paragraphs[paragraph_index])

composition_text = " ".join([paragraphs[18], *paragraphs[19:23]])
compare("Conseil - composition", board["composition"].replace("•", ""), composition_text)

committee_paragraphs = {
    "gouvernance-suivi": ((50, 51), range(53, 58)),
    "allocation-actifs": ((63, 64), range(66, 71)),
    "audit": ((76, 77), range(79, 84)),
    "pilotage-actuariel": ((89, 90), range(92, 97)),
    "risques": ((101, 102), range(104, 109)),
}
committees = {committee["id"]: committee for committee in system["governanceCommittees"]}
compare("Comités - introduction", system["governanceCommitteesIntroduction"], paragraphs[33])
for committee_id, (description_indexes, mission_indexes) in committee_paragraphs.items():
    committee = committees[committee_id]
    compare(f"{committee_id} - description", committee["description"], " ".join(paragraphs[index] for index in description_indexes))
    for index, paragraph_index in enumerate(mission_indexes):
        compare(f"{committee_id} - mission {index + 1}", committee["missions"][index], paragraphs[paragraph_index])

table_targets = [
    *board["currentComposition"],
    committees["gouvernance-suivi"],
    committees["allocation-actifs"],
    committees["audit"],
    committees["pilotage-actuariel"],
    committees["risques"],
]
for table_index, target in enumerate(table_targets):
    members = target["members"]
    rows = document.tables[table_index].rows[1:]
    if len(members) != len(rows):
        differences.append((f"Tableau {table_index} - nombre de membres", str(len(members)), str(len(rows))))
        continue
    for row_index, (member, row) in enumerate(zip(members, rows), start=1):
        compare(f"Tableau {table_index} ligne {row_index} - membre", member["name"], row.cells[0].text)
        compare(f"Tableau {table_index} ligne {row_index} - qualité", member["role"], row.cells[1].text)

compare("Missions - introduction", missions_data["governanceMissions"], paragraphs[115])
mission_items = missions_data["governanceMissionItems"]
compare("Mission 1 - titre", mission_items[0]["title"], paragraphs[116].split("–", 1)[1])
compare("Mission 1 - description", mission_items[0]["description"], paragraphs[117])
for index, paragraph_index in enumerate(range(118, 123)):
    compare(f"Mission 1 - élément {index + 1}", mission_items[0]["items"][index], paragraphs[paragraph_index])
compare("Mission 2 - titre", mission_items[1]["title"], paragraphs[123].split("–", 1)[1])
compare("Mission 2 - description", mission_items[1]["description"], paragraphs[124])
compare("Mission 3 - titre", mission_items[2]["title"], paragraphs[125].split("–", 1)[1])
compare("Mission 3 - description", mission_items[2]["description"], paragraphs[126])

if differences:
    print(f"DIFFERENCES: {len(differences)}")
    for label, site, word in differences:
        print(f"\n[{label}]\nSITE: {site}\nWORD: {word}")
else:
    print("OK: tous les textes métier et tableaux comparés correspondent au document Word.")
